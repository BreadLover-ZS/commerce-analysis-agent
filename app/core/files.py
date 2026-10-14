import io
import json
import multiprocessing
import os
import re
import sys
import zipfile
from pathlib import Path


def safe_path(root: Path, relative: str) -> Path:
    if not relative or "\\" in relative or "\x00" in relative:
        raise ValueError("无效文件路径")
    path = (root.resolve() / relative).resolve()
    if Path(relative).is_absolute() or not path.is_relative_to(root.resolve()):
        raise ValueError("文件路径越界")
    return path


def valid_filename(name: str) -> str:
    if len(name) > 120 or name.startswith(".") or re.search(r"[/\\\x00-\x1f]", name):
        raise ValueError("文件名不允许包含目录或控制字符")
    if Path(name).suffix.lower() not in {".md", ".txt", ".pdf", ".docx", ".xlsx"}:
        raise ValueError("支持 Markdown、TXT、PDF、DOCX、XLSX")
    return name


def _parse_upload(name: str, content: bytes) -> list[tuple[str, str]]:
    valid_filename(name)
    ext = Path(name).suffix.lower()
    if ext in {".docx", ".xlsx"}:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if (
                sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024
                or len(archive.infolist()) > 2000
            ):
                raise ValueError("压缩文档展开后过大")
    if ext in {".md", ".txt"}:
        text = content.decode("utf-8")
        fragments = [(f"lines=1-{len(text.splitlines())}", text)]
    elif ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted or len(reader.pages) > 100:
            raise ValueError("不支持加密或超过100页的PDF")
        fragments = [(f"page={i + 1}", p.extract_text() or "") for i, p in enumerate(reader.pages)]
    elif ext == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs]
        paragraphs += [
            " | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows
        ]
        fragments = [("paragraphs-and-tables", "\n".join(paragraphs))]
    else:
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            fragments = []
            for sheet in book.worksheets[:5]:
                if (
                    sheet.max_row
                    and sheet.max_row > 5000
                    or sheet.max_column
                    and sheet.max_column > 50
                ):
                    raise ValueError("表格超过5000行或50列")
                rows = [
                    " | ".join(str(c) if c is not None else "" for c in row)
                    for row in sheet.iter_rows(values_only=True)
                ]
                fragments.append((f"sheet={sheet.title}", "\n".join(rows)))
        finally:
            book.close()
    total = sum(len(text) for _, text in fragments)
    if total > 200000 or not any(text.strip() for _, text in fragments):
        raise ValueError("文档文本为空或超过20万字符；扫描PDF需要先进行OCR")
    return fragments


def _parse_process(sender, name, content):
    # Resource isolation, not an OS security sandbox. Libraries never get model keys.
    os.environ.clear()
    if sys.platform != "win32":
        import resource

        resource.setrlimit(resource.RLIMIT_CPU, (8, 8))
        if sys.platform == "linux":
            resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
    try:
        result = {"fragments": _parse_upload(name, content)}
    except Exception:
        result = {"error": "附件解析失败或超过资源限制"}
    sender.send_bytes(json.dumps(result, ensure_ascii=False).encode())
    sender.close()


def parse_upload(name: str, content: bytes) -> list[tuple[str, str]]:
    valid_filename(name)
    if len(content) > 20 * 1024 * 1024:
        raise ValueError("附件过大")
    context = multiprocessing.get_context("spawn")
    receiver, sender = context.Pipe(duplex=False)
    process = context.Process(target=_parse_process, args=(sender, name, content), daemon=True)
    process.start()
    sender.close()
    try:
        if not receiver.poll(12):
            raise ValueError("附件解析超时")
        # JSON over a bounded pipe avoids unpickling a parser-controlled object.
        result = json.loads(receiver.recv_bytes(maxlength=1_000_000))
        if "error" in result:
            raise ValueError(result["error"])
        return [(location, text) for location, text in result["fragments"]]
    finally:
        receiver.close()
        if process.is_alive():
            process.terminate()
        process.join(timeout=2)
        if process.is_alive():
            process.kill()
            process.join()
